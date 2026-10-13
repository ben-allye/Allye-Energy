"""Build an editable, ALLYE-branded PowerPoint from the HTML deck's data."""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import tempfile
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import cairosvg
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
HTML = ROOT / "notebooks/throughput_summary_deck.html"
OUT = ROOT / "reports/throughput_summary_allye.pptx"
LEGACY_OUT = ROOT / "notebooks/throughput_summary_allye.pptx"
REVIEW_OUT = ROOT / "reports/throughput_summary_allye_rev2.pptx"
LOGO = ROOT / "assets/allye_logo.png"
SVG_DIR = ROOT / "reports/throughput_svgs"

raw = HTML.read_text()
DATA = json.loads(re.search(r"const DATA = (\{.*?\});", raw, re.S).group(1))
SUMMARY = DATA["summary"]
START = datetime.fromisoformat(SUMMARY["start_date"])
END = datetime.fromisoformat(SUMMARY["end_date"])
PERIOD = f"{START:%-d %B} to {END:%-d %B %Y}"
PERIOD_SHORT = f"{START:%-d %b}–{END:%d %b %Y}"
UNITS = DATA["units"]
EFC_VALUES = DATA["fleet_daily_efc_values"]
EFC_MEAN = sum(EFC_VALUES) / len(EFC_VALUES)
EFC_MEDIAN = sorted(EFC_VALUES)[len(EFC_VALUES)//2] if len(EFC_VALUES)%2 else sum(sorted(EFC_VALUES)[len(EFC_VALUES)//2-1:len(EFC_VALUES)//2+1])/2
EFC_MAX = max(EFC_VALUES)
EFF_VALUES = [DATA["efficiency"][ep]["median"]*100 for ep in UNITS]
EFF_MIN, EFF_MAX = min(EFF_VALUES), max(EFF_VALUES)
WEEKDAY = DATA["weekday_kwh"]
WEEKDAY_TOP = sorted(["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"], key=lambda d: WEEKDAY[d], reverse=True)[:2]
PEAK_DIS_HOUR = max(range(24), key=lambda h: DATA["tou_hourly_out"][h])
PEAK_CHG_HOUR = max(range(24), key=lambda h: DATA["tou_hourly_in"][h])
CHARGE_FASTEST = min(UNITS, key=lambda ep: DATA["charge_time"][ep]["median_h"])
OUTLIER_UNIT = max(UNITS, key=lambda ep: DATA["box_stats"][ep]["max"])
OTHER_MAX = max(DATA["box_stats"][ep]["max"] for ep in UNITS if ep != OUTLIER_UNIT)
MONTHLY = list(zip(DATA["months"], DATA["fleet_monthly"]))
FULL_MONTHS = MONTHLY[1:-1] if len(MONTHLY) > 2 else MONTHLY
PEAK_MONTH, PEAK_MONTH_KWH = max(FULL_MONTHS, key=lambda row: row[1])
PEAK_MONTH_NAME = datetime.fromisoformat(PEAK_MONTH+"-01").strftime("%B")

ORANGE = "EB5F36"
PINK = "E7AECF"
BLUE = "78BEDA"
PURPLE = "514898"
GREEN = "4D6A28"
BROWN = "493A3A"
BEIGE = "E0C8B1"
WHITE = "FFFFFF"
PALE = "F8F5F0"
GRID = "E6DCD2"
MUTED = "6E6262"
UNIT = [ORANGE, BLUE, PURPLE, GREEN]

SVG_PICTURES = []


def prepare_svg_assets():
    """Capture the HTML deck's generated SVGs and add explicit SVG styling."""
    subprocess.run(["node", str(ROOT / "scripts/export_throughput_svgs.js")], check=True, cwd=ROOT)
    palette = {f"var(--u{i+1})": f"#{color}" for i, color in enumerate(UNIT)}
    palette.update({
        "var(--accent)": f"#{ORANGE}",
        "var(--panel-2)": f"#{WHITE}",
        "var(--ink)": f"#{BROWN}",
        "var(--ink-faint)": f"#{MUTED}",
        "var(--ink-soft)": f"#{MUTED}",
        "var(--line)": f"#{GRID}",
    })
    result = {}
    chart_names = ["monthly", "per_live", "weekday", "tou_hourly", "tou_heatmap",
                   "efc_line", "efficiency", "charge_time", "boxplot"]
    chart_names += [f"histogram_{i}" for i in range(1, 5)]
    for name in chart_names:
        raw = (SVG_DIR / f"{name}.raw.svg").read_text()
        raw = raw.replace("<svg ", '<svg xmlns="http://www.w3.org/2000/svg" ', 1)
        root = etree.fromstring(raw.encode())
        view = root.get("viewBox").split()
        root.set("width", view[2])
        root.set("height", view[3])
        root.attrib.pop("style", None)
        for elem in list(root.iter()):
            cls = elem.get("class", "")
            if cls in ("hit", "crosshair") or "display:none" in elem.get("style", ""):
                elem.getparent().remove(elem)
                continue
            for key, value in list(elem.attrib.items()):
                for old, new in palette.items():
                    value = value.replace(old, new)
                if name == "per_live":
                    value = value.replace(f"#{ORANGE}", f"#{BLUE}")
                elem.set(key, value)
            if cls in ("grid-line", "axis-line"):
                elem.set("stroke", f"#{GRID}")
                elem.set("stroke-width", "1")
            elif cls == "tick-label":
                elem.set("fill", f"#{MUTED}")
                elem.set("font-family", "Manrope")
                elem.set("font-size", "10.5")
            elif cls == "value-label":
                elem.set("fill", f"#{BROWN}")
                elem.set("font-family", "Manrope")
                elem.set("font-size", "11")
                elem.set("font-weight", "700")
            elif cls == "go-live-line":
                elem.set("stroke", f"#{MUTED}")
                elem.set("stroke-width", "1")
                elem.set("stroke-dasharray", "3,3")
                elem.set("opacity", "0.7")
            elif cls == "facet-title":
                elem.set("font-family", "Manrope")
                elem.set("font-size", "12")
                elem.set("font-weight", "700")
                style = elem.attrib.pop("style", "")
                match = re.search(r"fill:([^;]+)", style)
                if match:
                    elem.set("fill", match.group(1))
            elem.attrib.pop("class", None)
        svg_path = SVG_DIR / f"{name}.svg"
        png_path = SVG_DIR / f"{name}.png"
        data = etree.tostring(root, encoding="utf-8", xml_declaration=True)
        svg_path.write_bytes(data)
        cairosvg.svg2png(bytestring=data, write_to=str(png_path), scale=2)
        result[name] = (svg_path, png_path)
    return result


SVG_ASSETS = prepare_svg_assets()

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def C(hexstr):
    return RGBColor.from_string(hexstr)


def purple_heat(value, maximum):
    """Continuous single-hue scale; sqrt spacing keeps moderate hours visible."""
    low, high = "F3F1FA", PURPLE
    t = (max(0, value) / maximum) ** 0.5 if maximum else 0
    return "".join(f"{round(int(low[i:i+2], 16) * (1-t) + int(high[i:i+2], 16) * t):02X}"
                   for i in (0, 2, 4))


def box(slide, x, y, w, h, fill=None, line=None, radius=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        sh.adjustments[0] = 0.08
    sh.fill.solid()
    sh.fill.fore_color.rgb = C(fill or WHITE)
    sh.line.fill.background() if line is None else None
    if line:
        sh.line.color.rgb = C(line)
        sh.line.width = Pt(0.7)
    return sh


def txt(slide, text, x, y, w, h, size=14, color=BROWN, bold=False,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, font="Manrope"):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.01)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.text = str(text)
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = C(color)
    return sh


def line(slide, x1, y1, x2, y2, color=GRID, width=1):
    sh = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = C(color)
    sh.line.width = Pt(width)
    return sh


def base(section, title, n, dark=False):
    s = prs.slides.add_slide(blank)
    bg = PURPLE if dark else PALE
    box(s, 0, 0, 13.333, 7.5, bg)
    box(s, 0, 0, 0.14, 7.5, ORANGE)
    txt(s, section.upper(), 0.55, 0.28, 7.5, 0.28, 10, PINK if dark else PURPLE, True)
    txt(s, title, 0.55, 0.72, 11.95, 0.72, 27, WHITE if dark else BROWN,
        font="Aether")
    if LOGO.exists():
        # The reference deck's official wordmark, extracted without alteration.
        box(s, 11.72, 0.22, 1.17, 0.39, WHITE, None, True)
        s.shapes.add_picture(str(LOGO), Inches(11.86), Inches(0.30), width=Inches(0.89))
    line(s, 0.55, 7.12, 12.8, 7.12, PINK if dark else GRID, 0.7)
    txt(s, "ALLYE  /  SITECHARGE FLEET", 0.55, 7.17, 5.0, 0.18, 8, PINK if dark else MUTED, True)
    txt(s, f"{n:02d} / 09", 11.5, 7.17, 1.3, 0.18, 8, PINK if dark else MUTED, True, PP_ALIGN.RIGHT)
    return s


def stat(s, x, y, w, value, label, accent=ORANGE, height=1.05):
    box(s, x, y, w, height, WHITE, GRID, True)
    box(s, x, y, 0.075, height, accent)
    txt(s, value, x + 0.18, y + 0.12, w - 0.27, 0.43, 23, accent, True)
    txt(s, label, x + 0.18, y + 0.61, w - 0.27, 0.28, 10.2, MUTED)


def note(s, text, y=6.68):
    txt(s, text, 0.62, y, 12.0, 0.32, 9.2, MUTED)


def chart_frame(s, x, y, w, h, label):
    box(s, x, y, w, h, WHITE, GRID, True)
    txt(s, label.upper(), x + 0.2, y + 0.14, w - 0.4, 0.22, 10, PURPLE, True)


def bars(s, values, labels, x, y, w, h, ymax, color=ORANGE, fmt=lambda v: f"{v:,.0f}",
         colors=None, baseline=True, value_labels=True, label_size=9):
    left, right = x + 0.35, x + w - 0.12
    top, bottom = y + 0.13, y + h - 0.38
    for j in range(4):
        yy = bottom - (bottom - top) * j / 3
        line(s, left, yy, right, yy, GRID, 0.5)
    count = len(values)
    step = (right - left) / count
    bw = min(step * 0.62, 0.5)
    for i, (v, lab) in enumerate(zip(values, labels)):
        bx = left + step * i + (step - bw) / 2
        bh = max(0.015, (bottom - top) * max(v, 0) / ymax)
        box(s, bx, bottom - bh, bw, bh, colors[i] if colors else color)
        if value_labels:
            txt(s, fmt(v), bx - 0.25, bottom - bh - 0.28, bw + 0.5, 0.22, 8.5, BROWN, True, PP_ALIGN.CENTER)
        txt(s, lab, bx - 0.23, bottom + 0.06, bw + 0.46, 0.22, label_size, MUTED, False, PP_ALIGN.CENTER)


def native_chart(slide, kind, categories, series, x, y, w, h, ymax, major_unit,
                 point_colors=None, value_labels=False, number_format="0.0"):
    """Add a chart with its own editable PowerPoint data workbook."""
    chart_data = CategoryChartData()
    chart_data.categories = categories
    for label, values, _ in series:
        chart_data.add_series(label, values)
    chart = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), chart_data).chart
    chart.has_legend = False
    chart.has_title = False
    chart.chart_style = 10
    va = chart.value_axis
    va.minimum_scale = 0
    va.maximum_scale = ymax
    va.major_unit = major_unit
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = C(GRID)
    va.major_gridlines.format.line.width = Pt(0.6)
    va.tick_labels.font.name = "Manrope"
    va.tick_labels.font.size = Pt(8)
    va.tick_labels.font.color.rgb = C(MUTED)
    va.tick_labels.number_format = number_format
    ca = chart.category_axis
    ca.tick_labels.font.name = "Manrope"
    ca.tick_labels.font.size = Pt(8)
    ca.tick_labels.font.color.rgb = C(MUTED)
    ca.tick_label_position = XL_TICK_LABEL_POSITION.NEXT_TO_AXIS
    for si, (_, _, color) in enumerate(series):
        ser = chart.series[si]
        if kind == XL_CHART_TYPE.COLUMN_CLUSTERED:
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = C(color)
            ser.format.line.fill.background()
            if point_colors:
                for point, point_color in zip(ser.points, point_colors):
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = C(point_color)
        else:
            ser.format.line.color.rgb = C(color)
            ser.format.line.width = Pt(2.2)
    if kind == XL_CHART_TYPE.COLUMN_CLUSTERED:
        chart.plots[0].gap_width = 65
        if value_labels:
            plot = chart.plots[0]
            plot.has_data_labels = True
            plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
            plot.data_labels.font.name = "Manrope"
            plot.data_labels.font.size = Pt(8)
            plot.data_labels.font.color.rgb = C(BROWN)
            plot.data_labels.number_format = number_format
    return chart


def native_histogram(slide, counts, color, x, y, w, h):
    """Compact PowerPoint chart with editable bin counts in its workbook."""
    data = CategoryChartData()
    data.categories = [f"{i/10:.1f}" for i in range(len(counts))]
    data.add_series("Days", counts)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,
                                   Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    chart.has_legend = False
    chart.has_title = False
    chart.category_axis.visible = False
    chart.value_axis.visible = False
    chart.category_axis.format.line.color.rgb = C(GRID)
    chart.value_axis.format.line.fill.background()
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = DATA["hist_max_count"] + 2
    chart.plots[0].gap_width = 25
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = C(color)
    chart.series[0].format.line.fill.background()
    return chart


def svg_picture(slide, name, x, y, w, h):
    """Place a PNG fallback, then register its SVG for OOXML vector embedding."""
    svg_path, png_path = SVG_ASSETS[name]
    picture = slide.shapes.add_picture(str(png_path), Inches(x), Inches(y), Inches(w), Inches(h))
    SVG_PICTURES.append((len(prs.slides), picture.shape_id, svg_path))
    return picture


def embed_svg_parts(package):
    """Attach SVG to each picture using PowerPoint's svgBlip extension."""
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
    REL = "http://schemas.openxmlformats.org/package/2006/relationships"
    CT = "http://schemas.openxmlformats.org/package/2006/content-types"
    with ZipFile(package) as src:
        parts = {info.filename: src.read(info.filename) for info in src.infolist()}
        infos = src.infolist()
    ct_root = etree.fromstring(parts["[Content_Types].xml"])
    if not any(el.get("Extension") == "svg" for el in ct_root):
        etree.SubElement(ct_root, f"{{{CT}}}Default", Extension="svg", ContentType="image/svg+xml")
    parts["[Content_Types].xml"] = etree.tostring(ct_root, encoding="utf-8", xml_declaration=True)
    for slide_num, shape_id, svg_path in SVG_PICTURES:
        slide_part = f"ppt/slides/slide{slide_num}.xml"
        rels_part = f"ppt/slides/_rels/slide{slide_num}.xml.rels"
        slide_root = etree.fromstring(parts[slide_part])
        rels_root = etree.fromstring(parts[rels_part])
        used = [int(el.get("Id")[3:]) for el in rels_root if re.fullmatch(r"rId\d+", el.get("Id", ""))]
        rid = f"rId{max(used, default=0)+1}"
        media_name = f"vector_{svg_path.stem}.svg"
        etree.SubElement(rels_root, f"{{{REL}}}Relationship", Id=rid,
                         Type=f"{R}/image", Target=f"../media/{media_name}")
        pic = next((el for el in slide_root.findall(f".//{{{P}}}pic")
                    if (cnv := el.find(f"./{{{P}}}nvPicPr/{{{P}}}cNvPr")) is not None
                    and cnv.get("id") == str(shape_id)), None)
        if pic is None:
            raise RuntimeError(f"Picture shape {shape_id} missing on slide {slide_num}")
        blip = pic.find(f".//{{{A}}}blip")
        ext_lst = etree.SubElement(blip, f"{{{A}}}extLst")
        ext = etree.SubElement(ext_lst, f"{{{A}}}ext",
                               uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
        etree.SubElement(ext, f"{{{ASVG}}}svgBlip", nsmap={"asvg": ASVG}).set(f"{{{R}}}embed", rid)
        parts[slide_part] = etree.tostring(slide_root, encoding="utf-8", xml_declaration=True)
        parts[rels_part] = etree.tostring(rels_root, encoding="utf-8", xml_declaration=True)
        parts[f"ppt/media/{media_name}"] = svg_path.read_bytes()
    fd, temp_name = tempfile.mkstemp(suffix=".pptx", dir=package.parent)
    os.close(fd)
    try:
        with ZipFile(temp_name, "w") as out:
            for info in infos:
                out.writestr(info, parts.pop(info.filename))
            for name, data in parts.items():
                out.writestr(name, data, compress_type=ZIP_DEFLATED)
        os.replace(temp_name, package)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)


def small_table(s, x, y, widths, headers, rows, row_h=0.45, size=11):
    total = sum(widths)
    box(s, x, y, total, 0.42, PURPLE)
    xx = x
    for width, head in zip(widths, headers):
        txt(s, head, xx + 0.12, y + 0.04, width - 0.17, 0.31, 9.2, WHITE, True)
        xx += width
    for ri, row in enumerate(rows):
        yy = y + 0.42 + ri * row_h
        box(s, x, yy, total, row_h, WHITE if ri % 2 == 0 else PALE)
        xx = x
        for width, cell in zip(widths, row):
            txt(s, cell, xx + 0.12, yy + 0.03, width - 0.17, row_h - 0.06, size, BROWN, ri == 0 and False)
            xx += width


# 1 — title and scope
s = base("Throughput summary", "Sitecharge fleet throughput", 1, True)
txt(s, PERIOD, 0.59, 1.55, 9.8, 0.43, 19, WHITE)
txt(s, f"{len(UNITS)} Sitecharge units  ·  {SUMMARY['total_cycles']:,} measured cycles", 0.59, 2.08, 10.6, 0.4, 15, PINK)
box(s, 0.59, 2.92, 12.1, 2.45, WHITE, None, True)
txt(s, "What this deck covers", 0.88, 3.2, 4.4, 0.35, 16, PURPLE, font="Aether")
for i, (v, lab) in enumerate([
    (f"{SUMMARY['total_throughput_kwh']/1000:,.2f} MWh", "fleet throughput"),
    (f"{SUMMARY['fleet_charge_median_h']:.2f} h", "median full-charge equivalent"),
    (f"{EFF_MIN:.1f}–{EFF_MAX:.1f}%", "unit median round-trip efficiency")]):
    x = 0.91 + 4.02 * i
    if i: line(s, x - 0.25, 3.72, x - 0.25, 4.94, GRID)
    txt(s, v, x, 3.79, 3.75, 0.51, 26, [ORANGE, BLUE, GREEN][i], True)
    txt(s, lab, x, 4.42, 3.75, 0.35, 11, MUTED)
txt(s, "Throughput, weekly use, time of use, daily cycling, efficiency, charge time, and unit distribution.",
    0.6, 5.77, 12.0, 0.48, 15, WHITE)

# 2 — monthly throughput
s = base("Fleet throughput", "Monthly throughput across the live fleet", 2)
chart_frame(s, 0.6, 1.68, 6.0, 4.3, "Fleet throughput  ·  MWh per month")
chart_frame(s, 6.77, 1.68, 5.95, 4.3, "Throughput per live unit  ·  MWh per month")
svg_picture(s,"monthly",0.77,2.35,5.65,2.94)
svg_picture(s,"per_live",6.93,2.35,5.62,2.93)
note(s, f"{START:%B} and {END:%B} are partial months. {PEAK_MONTH_NAME} is the highest full month at {PEAK_MONTH_KWH/1000:.1f} MWh.", 6.33)

# 3 — weekly rhythm
s = base("Usage rhythm", f"{WEEKDAY_TOP[0]} and {WEEKDAY_TOP[1]} carry the most throughput", 3)
chart_frame(s, 0.6, 1.68, 7.65, 4.67, "Average fleet throughput  ·  kWh per day")
svg_picture(s,"weekday",0.87,2.14,7.08,3.85)
stat(s, 8.55, 1.68, 4.13, f"{WEEKDAY[WEEKDAY_TOP[0]]:.0f} kWh", f"{WEEKDAY_TOP[0]} fleet average", ORANGE, 1.25)
stat(s, 8.55, 3.1, 4.13, f"{WEEKDAY['Saturday']:.0f} kWh", "Saturday fleet average", PINK, 1.25)
stat(s, 8.55, 4.52, 4.13, f"{WEEKDAY['Sunday']:.0f} kWh", "Sunday fleet average", GREEN, 1.25)
note(s, f"Saturday is {SUMMARY['saturday_pct_weekday']}% of the Monday–Friday average. Values are daily fleet averages over the recorded period.", 6.65)

# 4 — hourly charge/discharge and heatmap
s = base("Time of use", f"Discharge peaks at {PEAK_DIS_HOUR:02d}:00; charge peaks at {PEAK_CHG_HOUR:02d}:00", 4)
chart_frame(s, 0.6, 1.62, 12.08, 3.0, "Average weekday energy  ·  kWh by hour")
svg_picture(s,"tou_hourly",1.53,1.94,10.1,2.63)
box(s, 9.15, 1.79, 0.18, 0.08, ORANGE)
txt(s, "Discharge", 9.42, 1.72, 1.18, 0.22, 9, MUTED)
box(s, 10.67, 1.79, 0.18, 0.08, BLUE)
txt(s, "Charge", 10.94, 1.72, 1.0, 0.22, 9, MUTED)
chart_frame(s, 0.6, 4.74, 12.08, 2.12, "Discharge heatmap  ·  total kWh by weekday and hour (square-root colour scale)")
svg_picture(s,"tou_heatmap",1.98,5.08,9.37,1.78)
txt(s, "0", 9.32, 6.91, 0.23, 0.13, 7.4, MUTED, align=PP_ALIGN.RIGHT)
for i in range(16):
    box(s, 9.62+i*0.117, 6.93, 0.114, 0.07, purple_heat(DATA["tou_heatmap_max"]*i/15, DATA["tou_heatmap_max"]))
txt(s, f"{DATA['tou_heatmap_max']/1000:.1f} MWh", 11.58, 6.90, 0.88, 0.14, 7.4, MUTED)

# 5 — daily cycles
s = base("Daily cycling", f"Median fleet use is {EFC_MEDIAN:.2f} equivalent cycles a day", 5)
chart_frame(s, 0.6, 1.64, 12.08, 3.75, "Fleet equivalent full cycles  ·  daily")
svg_picture(s,"efc_line",0.93,2.12,11.26,2.93)
stat(s,0.6,5.61,3.87,f"{EFC_MEDIAN:.2f}","median fleet EFC / day",PURPLE,0.91)
stat(s,4.73,5.61,3.87,f"{EFC_MEAN:.2f}","mean fleet EFC / day",BLUE,0.91)
stat(s,8.86,5.61,3.82,f"{EFC_MAX:.2f}","busiest day, fleet EFC",ORANGE,0.91)
note(s,f"EFC uses pack-outage-corrected capacity. A fleet total of {len(UNITS):.1f} equals one full cycle per unit.",6.7)

# 6 — efficiency
s=base("Round-trip efficiency",f"Unit medians span {EFF_MIN:.1f}% to {EFF_MAX:.1f}%",6)
chart_frame(s,0.6,1.68,7.2,4.74,"Median efficiency with p10–p90 range")
eff=DATA["efficiency"]
svg_picture(s,"efficiency",0.93,2.2,6.53,3.97)
rows=[(ep[-4:],str(eff[ep]["n"]),f'{eff[ep]["median"]*100:.1f}%',
       f'{eff[ep]["p10"]*100:.1f}–{eff[ep]["p90"]*100:.1f}%') for ep in DATA["units"]]
small_table(s,8.04,1.68,[0.77,0.8,0.92,2.15],["Unit","Cycles","Median","p10–p90"],rows,0.58,10.5)
stat(s,8.04,4.79,4.64,f"{SUMMARY['total_cycles']:,}","cycles with >5 kWh in and out",ORANGE,1.14)
note(s,"Efficiency = discharge energy ÷ charge energy per cycle. The bars show each unit's 10th to 90th percentile.",6.7)

# 7 — charge time
s=base("Charge time",f"Fleet median full-charge equivalent is {SUMMARY['fleet_charge_median_h']:.2f} hours",7)
chart_frame(s,0.6,1.68,7.2,4.76,"Median equivalent full-charge time  ·  hours")
charge=DATA["charge_time"]
svg_picture(s,"charge_time",0.95,2.22,6.35,3.59)
rows=[(ep[-4:],str(charge[ep]["n"]),f'{charge[ep]["median_h"]:.2f}',
       f'{charge[ep]["mean_h"]:.2f}') for ep in DATA["units"]]
small_table(s,8.03,1.68,[0.8,1.05,1.4,1.39],["Unit","Sessions","Median h","Mean h"],rows,0.58,10.5)
stat(s,8.03,4.8,4.65,f"{charge[CHARGE_FASTEST]['median_h']:.2f} h",f"fastest unit: {CHARGE_FASTEST[-4:]}",GREEN,1.11)
note(s,"Each session's charge rate is projected to a 0–100% swing. Sessions with less than 5% SoC gain are excluded.",6.7)

# 8 — SVG boxplot and histograms
s=base("Unit distribution",f"Unit daily medians span {min(DATA['box_stats'][ep]['median'] for ep in UNITS):.2f}–{max(DATA['box_stats'][ep]['median'] for ep in UNITS):.2f} cycles",8)
chart_frame(s,0.6,1.68,6.05,4.68,"Daily EFC spread  ·  box / median / whiskers")
svg_picture(s,"boxplot",0.84,2.11,5.58,3.72)
chart_frame(s,6.82,1.68,5.86,4.68,"Daily EFC distribution by unit")
for i in range(4):
    x=7.05+(i%2)*2.72
    y=2.14+(i//2)*2.01
    svg_picture(s,f"histogram_{i+1}",x,y,2.62,1.93)
note(s,f"Unit {OUTLIER_UNIT[-4:]} reaches {DATA['box_stats'][OUTLIER_UNIT]['max']:.2f} EFC on its highest day; other unit maxima are at or below {OTHER_MAX:.2f}. Histogram values above 2.0 are clipped.",6.66)

# 9 — concise readout
s=base("Summary","Seven observations from the measured period",9,True)
items=[
    ("01","Fleet throughput",f"{PEAK_MONTH_NAME} is the highest full month at {PEAK_MONTH_KWH/1000:.1f} MWh."),
    ("02","Weekly rhythm",f"{WEEKDAY_TOP[0]} and {WEEKDAY_TOP[1]} average {WEEKDAY[WEEKDAY_TOP[0]]:.0f} and {WEEKDAY[WEEKDAY_TOP[1]]:.0f} kWh."),
    ("03","Daily timing",f"Discharge peaks at {PEAK_DIS_HOUR:02d}:00; charging peaks at {PEAK_CHG_HOUR:02d}:00 on weekdays."),
    ("04","Cycling",f"Median fleet use is {EFC_MEDIAN:.2f} EFC/day, with a {EFC_MAX:.2f} EFC maximum."),
    ("05","Efficiency",f"Unit median round-trip efficiency ranges from {EFF_MIN:.1f}% to {EFF_MAX:.1f}%."),
    ("06","Charge time",f"Median full-charge equivalent is {SUMMARY['fleet_charge_median_h']:.2f} hours across the fleet."),
    ("07","Outlier",f"Unit {OUTLIER_UNIT[-4:]} records a {DATA['box_stats'][OUTLIER_UNIT]['max']:.2f} EFC day; other maxima are ≤{OTHER_MAX:.2f}."),
]
for i,(num,head,body) in enumerate(items):
    col=0 if i<4 else 1; row=i if col==0 else i-4
    x=0.62+6.1*col; y=1.78+1.16*row
    box(s,x,y,5.83,1.00,WHITE,None,True)
    box(s,x,y,0.09,1.0,[ORANGE,PINK,BLUE,PURPLE,GREEN,BEIGE,ORANGE][i])
    txt(s,num,x+0.22,y+0.16,0.46,0.32,11,PURPLE,True)
    txt(s,head,x+0.76,y+0.10,4.76,0.28,13,BROWN,True)
    txt(s,body,x+0.76,y+0.42,4.78,0.42,10.6,MUTED)
txt(s,f"Period: {PERIOD_SHORT}. {START:%B} and {END:%B} are partial months.",6.76,5.56,5.72,0.52,11,PINK)

OUT.parent.mkdir(parents=True,exist_ok=True)
prs.save(OUT)
embed_svg_parts(OUT)
shutil.copy2(OUT, LEGACY_OUT)
shutil.copy2(OUT, REVIEW_OUT)
print(f"Wrote {OUT} ({len(prs.slides)} slides)")
